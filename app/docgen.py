"""Render a polished .docx from a plan plus its drafted sections."""
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def build_docx(plan, sections, assumptions, path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading(plan.get("title", "Generated Document"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        f"{plan.get('doc_type', 'document').replace('_', ' ').title()}  |  "
        f"Generated {date.today().isoformat()}"
    )
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if assumptions:
        doc.add_heading("Assumptions", level=1)
        for a in assumptions:
            doc.add_paragraph(a, style="List Bullet")

    for heading, body in sections:
        doc.add_heading(heading, level=1)
        for para in (p.strip() for p in body.split("\n") if p.strip()):
            doc.add_paragraph(para)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Generated autonomously by the FastAPI agent pipeline."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in footer.runs:
        r.font.size = Pt(8)

    doc.save(path)
    return path
